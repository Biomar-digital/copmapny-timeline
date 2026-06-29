"""
Parse an incoming Word policy into a simple, render-agnostic content model.

We deliberately keep the model tiny (heading levels, paragraphs, bullets and
tables) because that is everything the BioMar policy template expresses. The
mapping from Word styles to our block types lives here so the generator stays
purely about layout.
"""
from dataclasses import dataclass, field
from typing import List, Optional

import docx
from docx.document import Document as _Doc
from docx.table import Table as _Table
from docx.text.paragraph import Paragraph as _Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.oxml.ns import qn


# ---- block types ----------------------------------------------------------
@dataclass
class Heading:
    level: int          # 1 or 2
    text: str

@dataclass
class Body:
    text: str
    bold: bool = False
    runs: list = None        # optional [(text, bold)] for run-in bold labels

@dataclass
class Bullet:
    text: str

@dataclass
class TableBlock:
    rows: List[List[str]]
    header: bool = True

@dataclass
class ImageBlock:
    data: bytes             # raw image bytes (png/jpeg)
    width: float = 0.0      # intended display size in points (0 = natural)
    height: float = 0.0

@dataclass
class Columns:
    cols: list = field(default_factory=list)   # list of columns; each a list of blocks

@dataclass
class Policy:
    title: str
    year: str
    blocks: list = field(default_factory=list)
    # Back-cover "Version history / Owner and approver" card (optional: the card
    # is only drawn when at least one of the fields below is supplied).
    version: str = "Version 1:"
    approval_date: str = ""
    owner: str = ""
    approver: str = ""
    # Signatures page (penultimate) adoption statement.
    adopted_on: str = ""
    effective_on: str = ""
    signatures: bool = True   # include the board signatures page
    cover_year: bool = True   # show the year on the cover (some originals omit it)
    lead_title: bool = False  # repeat the title as a lead heading on page 1 of body
    body_size: float = 11     # body point size (e.g. the Code of Conduct uses 10)
    footer_note: str = ""     # small disclaimer line under the footer address
    version_date: str = ""    # date shown on the version line (defaults to approval date)
    has_signed: bool = False  # this unsigned render has a signed sibling (-> drop the card)


def _para_images(item):
    """Inline images in a paragraph -> [(bytes, width_pt, height_pt)]."""
    out = []
    p = item._p
    for blip in p.findall(".//" + qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        try:
            data = item.part.related_parts[rid].blob
        except KeyError:
            continue
        w = h = 0.0
        ext = p.find(".//" + qn("wp:extent"))
        if ext is not None:
            try:
                w = int(ext.get("cx")) / 12700.0   # EMU -> points
                h = int(ext.get("cy")) / 12700.0
            except (TypeError, ValueError):
                pass
        out.append((data, w, h))
    return out


def _is_brand_artwork(data: bytes) -> bool:
    """Some source .docx files bake the designed cover, back cover, section
    dividers ("Audit Committee Charter") and full-page logos in as images. The
    template redraws its own cover and furniture, so these must not be rendered
    inline (otherwise they appear as stray cover pages / oversized logos). They
    are all dominated by the BioMar navy brand colour, or are full-page portrait
    artwork — genuine content images are neither."""
    try:
        from io import BytesIO
        from PIL import Image
        im = Image.open(BytesIO(data)).convert("RGB")
    except Exception:
        return False
    w, h = im.size
    px = list(im.resize((16, 16)).getdata())
    n = len(px) or 1
    mr = sum(p[0] for p in px) / n
    mg = sum(p[1] for p in px) / n
    mb = sum(p[2] for p in px) / n
    navy = mb > mr and mb > 70 and mr < 95 and mg < 120
    portrait_page = h > w and 0.6 < w / h < 0.8 and min(w, h) > 400
    return navy or portrait_page


def _iter_block_items(parent):
    """Yield paragraphs and tables in document order."""
    body = parent.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield _Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield _Table(child, parent)


def _clean(text: str) -> str:
    # Word often encodes the clause number + a tab ("3.1\tText"). Collapse the
    # tab to a regular space so it flows as one justified paragraph.
    return text.replace("\t", " ").replace("​", "").strip()


def _cell_text(cell) -> str:
    """Join a table cell's paragraphs, marking 'List Paragraph' items with a
    leading bullet so the generator can render them as a bulleted list inside the
    cell (the source keeps these as real list items; flattening loses them)."""
    out = []
    for p in cell.paragraphs:
        t = _clean(p.text)
        if not t:
            continue
        style = (p.style.name if p.style else "").lower()
        out.append("• " + t if "list" in style else t)
    return "\n".join(out)


def _heading_like(text: str) -> bool:
    # A heading is short and does not end like a running sentence. Incoming
    # Word files are often mis-styled (whole paragraphs tagged Heading 1, or
    # body sentences tagged Heading 2), so we confirm headings by shape too.
    return len(text.split()) <= 14 and not text.rstrip().endswith((".", ":", ";"))


def _classify(text: str, style: str):
    """Return ('heading', level) | ('bullet', 0) | ('body', 0)."""
    s = (style or "").lower()
    if "list" in s:
        return "bullet", 0
    if "heading 2" in s or s == "subtitle":
        return ("heading", 2) if _heading_like(text) else ("body", 0)
    if "heading 1" in s or s in ("heading", "title"):
        return ("heading", 1) if _heading_like(text) else ("body", 0)
    return "body", 0


def _body_from_runs(item, text):
    """Build a Body, keeping per-run bold (a run-in bold label + plain text) when
    the paragraph mixes bold and regular; otherwise collapse to a bold flag."""
    rlist = [(r.text, bool(r.bold)) for r in item.runs if r.text and (r.text.strip() or r.text == " ")]
    anyb = any(b for _, b in rlist)
    allb = bool(rlist) and all(b for _, b in rlist)
    if anyb and not allb:
        return Body(text=text, runs=rlist)
    return Body(text=text, bold=allb)


def parse_docx(path: str, title: Optional[str] = None,
               year: Optional[str] = None, **meta) -> Policy:
    doc = docx.Document(path)
    blocks = []
    doc_title = None

    for item in _iter_block_items(doc):
        if isinstance(item, _Table):
            # A borderless 1-row, 2-column table (no "Grid" style) encodes a
            # two-column text layout (e.g. a definitions / references page).
            style_name = item.style.name if item.style else ""
            if "Grid" not in (style_name or "") and len(item.columns) == 2 and len(item.rows) == 1:
                cols = []
                for ci in range(2):
                    sub = []
                    for para in item.rows[0].cells[ci].paragraphs:
                        t = _clean(para.text)
                        if not t:
                            continue
                        kind, level = _classify(t, para.style.name if para.style else "Normal")
                        if kind == "heading":
                            sub.append(Heading(level=level, text=t))
                        elif kind == "bullet":
                            sub.append(Bullet(text=t))
                        else:
                            sub.append(_body_from_runs(para, t))
                    cols.append(sub)
                if any(cols):
                    blocks.append(Columns(cols=cols))
                continue
            rows = [[_cell_text(c) for c in row.cells] for row in item.rows]
            rows = [r for r in rows if any(r)]
            if rows:
                blocks.append(TableBlock(rows=rows))
            continue

        for data, iw, ih in _para_images(item):
            if _is_brand_artwork(data):
                continue
            blocks.append(ImageBlock(data=data, width=iw, height=ih))

        text = _clean(item.text)
        if not text:
            continue
        style = item.style.name if item.style else "Normal"

        # The document's own title line: capture it, keep it out of the body.
        if (style or "").lower() == "title":
            if doc_title is None:
                doc_title = text
            continue
        if title and text.strip().lower() == title.strip().lower():
            continue
        if len(text) < 4 and not text[0].isdigit():   # stray fragments ("Com")
            continue

        kind, level = _classify(text, style)
        if kind == "heading":
            blocks.append(Heading(level=level, text=text))
        elif kind == "bullet":
            blocks.append(Bullet(text=text))
        else:
            blocks.append(_body_from_runs(item, text))

    if title is None:
        title = doc_title or "Policy"
    if year is None:
        import datetime
        year = str(datetime.date.today().year)

    meta = {k: v for k, v in meta.items() if v is not None}
    return Policy(title=title, year=year, blocks=blocks, **meta)
