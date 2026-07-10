"""
Render the same parsed Policy model as an editable, BioMar-branded Word (.docx).

This mirrors the PDF (`generator.py`) as closely as Word allows: navy headings
and body, run-in bold labels, bullets, tables, two-column blocks, a title page
and the "Version history / Owner and approver" card, plus the address footer.
It will not be pixel-identical to the PDF (Word lays out columns and fonts
differently), but it is a faithful, editable copy of the same content.
"""
from io import BytesIO
import os
import re
import tempfile

from docx import Document
from docx.shared import Pt, RGBColor, Emu, Mm, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

import model as M
import brand as B

NAVY = RGBColor(0x1C, 0x40, 0x76)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x6B, 0x87, 0xA4)
FONT = "Avenir Next LT Pro"          # Word substitutes if the licensed font is absent
FOOTER = ("BioMar Group A/S · Kalkværksvej 16, 15. · 8000 Aarhus C · "
          "Denmark · Tel +45 86 20 49 70 · www.biomar.com")
EMU_PER_PT = 12700
PAGE_W_PT, PAGE_H_PT = 595.276, 841.890   # A4, matching the PDF

# Hang-indent support for "hanging_indent" policies (Articles of Association,
# Remuneration Policy): mirrors generator.py's CLAUSE_INDENT/_NUM/_LETTERED so
# the editable Word copy matches the PDF's article-style layout instead of
# flush-left wrapped lines.
_NUM = re.compile(r"^(\d+(?:\.\d+)*\.?)(\s+)(.*)$", re.S)
_LETTERED = re.compile(r"^\(?[a-z][.)]\s")
CLAUSE_INDENT = 35.4
LETTER_MARKER_INDENT = 53.4 - 42.6
LETTER_TEXT_INDENT = 71.4 - 42.6


def _set_cell_bg(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _no_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "none")
        e.set(qn("w:sz"), "0")
        borders.append(e)
    tblPr.append(borders)


def _run(p, text, *, bold=False, italic=False, size=11, color=NAVY, font=FONT):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.name = font
    return r


def _heading(doc, text, level, hang=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(13 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 4)
    p.paragraph_format.keep_with_next = True
    if hang:
        p.paragraph_format.left_indent = Pt(CLAUSE_INDENT)
        p.paragraph_format.first_line_indent = Pt(-CLAUSE_INDENT)
    _run(p, text, bold=True, size=14 if level == 1 else 12)


def _body(doc, blk, size=11, indent_mode=None):
    """indent_mode: None (flush), "hang" (numbered clause — wrap aligns under
    the clause text), or "uniform" (content of a short sub-heading above —
    same left position on every line, no hanging first line)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(10)
    if indent_mode == "hang":
        p.paragraph_format.left_indent = Pt(CLAUSE_INDENT)
        p.paragraph_format.first_line_indent = Pt(-CLAUSE_INDENT)
    elif indent_mode == "uniform":
        p.paragraph_format.left_indent = Pt(CLAUSE_INDENT)
        p.paragraph_format.first_line_indent = Pt(0)
    if blk.runs:                              # mixed-style runs (bold label, italic word, ...)
        for t, b, i in blk.runs:
            _run(p, t, bold=b, italic=i, size=size)
    else:
        _run(p, blk.text, bold=blk.bold, italic=getattr(blk, "italic", False), size=size)
    return p


def _bullet(doc, text, size=11, hang=False):
    # A lettered item ("a. ...") already carries its own marker in the text,
    # so it must NOT also get Word's automatic "List Bullet" glyph — same
    # double-marker bug the PDF generator guards against. Use a plain
    # paragraph with a manual hanging indent instead.
    if hang:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(LETTER_TEXT_INDENT)
        p.paragraph_format.first_line_indent = Pt(-(LETTER_TEXT_INDENT - LETTER_MARKER_INDENT))
    else:
        p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    _run(p, text, size=size)


def _table(doc, blk):
    rows = blk.rows
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    try:
        t.style = "Table Grid"
    except KeyError:
        pass
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(ncol):
            txt = row[ci] if ci < len(row) else ""
            cell = cells[ci]
            cell.paragraphs[0].text = ""
            header = blk.header and ri == 0
            if header:
                _set_cell_bg(cell, "1C4076")
            _run(cell.paragraphs[0], txt, bold=header, size=8.5,
                 color=WHITE if header else NAVY)


def _render_blocks(doc, blocks, size=11, hanging_indent=False):
    # Mirrors generator.py's _story(): tracks whether the previous block was a
    # short numbered sub-heading with no clause number of its own ("3.2
    # Incentive pay"), so the paragraph(s) right after it — its actual content
    # — get the same indent even though they don't start with a number.
    prev_subhead = False
    for blk in blocks:
        if isinstance(blk, M.Heading):
            _heading(doc, blk.text, blk.level, hang=hanging_indent)
            prev_subhead = False
        elif isinstance(blk, M.Bullet):
            is_lettered = bool(_LETTERED.match(blk.text.strip()))
            _bullet(doc, blk.text, size, hang=hanging_indent and is_lettered)
            prev_subhead = False
        elif isinstance(blk, M.TableBlock):
            _table(doc, blk)
            prev_subhead = False
        elif isinstance(blk, M.ImageBlock):
            try:
                w = Emu(int(blk.width * EMU_PER_PT)) if blk.width else None
                doc.add_picture(BytesIO(blk.data), width=w)
            except Exception:
                pass
            prev_subhead = False
        elif isinstance(blk, M.Columns):
            t = doc.add_table(rows=1, cols=max(len(blk.cols), 1))
            _no_table_borders(t)
            for ci, col in enumerate(blk.cols):
                cell = t.rows[0].cells[ci]
                cell.paragraphs[0].text = ""
                first = True
                for sub in col:
                    # reuse the same renderers, but into the cell
                    _render_into_cell(cell, sub, size, first)
                    first = False
            prev_subhead = False
        elif isinstance(blk, M.Body):
            if not hanging_indent:
                _body(doc, blk, size)
                continue
            is_numbered = bool(_NUM.match(blk.text.strip()))
            if is_numbered:
                _body(doc, blk, size, indent_mode="hang")
                stripped = blk.text.strip()
                prev_subhead = (len(stripped.split()) <= 8
                                and not stripped.rstrip().endswith((".", ":", ";")))
            elif prev_subhead:
                _body(doc, blk, size, indent_mode="uniform")
                # leave prev_subhead as-is: a sub-heading's content can span
                # several paragraphs, all needing the same indent.
            else:
                _body(doc, blk, size)
                prev_subhead = False


def _render_into_cell(cell, blk, size, first):
    if isinstance(blk, M.Heading):
        p = cell.add_paragraph() if not first else cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0 if first else 8)
        p.paragraph_format.space_after = Pt(4)
        _run(p, blk.text, bold=True, size=12 if blk.level == 1 else 10.5)
    elif isinstance(blk, M.Bullet):
        p = cell.add_paragraph(style="List Bullet")
        _run(p, blk.text, size=size - 1)
    elif isinstance(blk, M.Body):
        p = cell.add_paragraph() if not first else cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if blk.runs:
            for t, b, i in blk.runs:
                _run(p, t, bold=b, italic=i, size=size - 1)
        else:
            _run(p, blk.text, bold=blk.bold, italic=getattr(blk, "italic", False), size=size - 1)


def _zero_margins(section):
    section.page_width = Mm(210); section.page_height = Mm(297)
    for a in ("top_margin", "bottom_margin", "left_margin", "right_margin",
              "header_distance", "footer_distance"):
        setattr(section, a, Pt(0))


def _normal_margins(section):
    section.page_width = Mm(210); section.page_height = Mm(297)
    # Top margin clears the big page-anchored logo (bottom at ~147pt), so body
    # text begins below it just like the PDF content pages.
    section.top_margin = Pt(146)
    section.bottom_margin = Pt(B.MARGIN_BOTTOM)
    section.left_margin = Pt(B.MARGIN_L)
    section.right_margin = Pt(B.MARGIN_R)
    section.header_distance = Pt(34)
    section.footer_distance = Pt(20)


def _cover_image_page(doc, pdf_path):
    """Embed page 1 of the rendered PDF as a full-bleed cover image (identical
    branding: background, logo, light-blue title, year, address)."""
    try:
        import fitz
        d = fitz.open(pdf_path)
        pix = d[0].get_pixmap(dpi=200)
        tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
        pix.save(tmp.name)
        d.close()
    except Exception:
        return False
    sec = doc.sections[0]
    _zero_margins(sec)
    p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(tmp.name, width=Mm(210))
    try:
        os.unlink(tmp.name)
    except OSError:
        pass
    return True


def _float_image(run, image_path, x_pt, y_pt, w_pt, h_pt):
    """Place an image as a page-anchored floating picture (top-left origin),
    matching the PDF's absolute logo position."""
    run.add_picture(image_path, width=Emu(_e(w_pt)), height=Emu(_e(h_pt)))
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    extent = inline.find(qn("wp:extent"))
    cx, cy = extent.get("cx"), extent.get("cy")
    graphic = inline.find(qn("a:graphic"))
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "r", "pic")} behindDoc="0" distT="0" distB="0" '
        f'distL="0" distR="0" simplePos="0" locked="0" layoutInCell="1" allowOverlap="1" '
        f'relativeHeight="2"><wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{_e(x_pt)}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{_e(y_pt)}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:wrapNone/>'
        f'<wp:docPr id="201" name="BioMar logo"/></wp:anchor>')
    anchor.append(graphic)
    drawing.replace(inline, anchor)


def _e(pt):
    return int(round(pt * EMU_PER_PT))


def _content_furniture(section, policy):
    """Running header (BioMar Group + title top-left, big logo top-right) and a
    footer with a thin rule above the company address — matching the PDF exactly."""
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    h = section.header.paragraphs[0]
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h.paragraph_format.space_after = Pt(0)
    _run(h, "BioMar Group", size=8, color=NAVY)
    h.add_run().add_break()
    _run(h, policy.title, size=8, color=NAVY)
    # Big logo, page-anchored top-right at the PDF's coordinates (136x130pt,
    # 16.6pt from the right edge, 16.7pt from the top).
    if os.path.exists(B.LOGO):
        _float_image(h.add_run(), B.LOGO,
                     x_pt=PAGE_W_PT - B.LOGO_W - 16.6, y_pt=B.LOGO_TOP,
                     w_pt=B.LOGO_W, h_pt=B.LOGO_H)
    f = section.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = f._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr"); top = OxmlElement("w:top")
    top.set(qn("w:val"), "single"); top.set(qn("w:sz"), "4")
    top.set(qn("w:space"), "6"); top.set(qn("w:color"), "C3E4EF")
    pbdr.append(top); pPr.append(pbdr)
    _run(f, FOOTER, size=8, color=NAVY)


def _version_card(doc, policy):
    if not any([policy.approval_date, policy.owner, policy.approver]):
        return
    if getattr(policy, "has_signed", False) and not policy.signatures:
        return  # unsigned sibling of a signed doc: no card (matches the PDF)
    doc.add_paragraph()
    t = doc.add_table(rows=1, cols=2)
    try:
        t.style = "Table Grid"
    except KeyError:
        pass
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = t.rows[0].cells
    _set_cell_bg(left, "1C4076"); _set_cell_bg(right, "1C4076")
    _run(left.paragraphs[0], "Version history", bold=True, size=10, color=WHITE)
    _run(right.paragraphs[0], "Owner and approver", bold=True, size=10, color=WHITE)
    body = t.add_row().cells
    vdate = policy.approval_date or getattr(policy, "version_date", "") or "—"
    lp = body[0].paragraphs[0]
    _run(lp, f"{policy.version} ", bold=True, size=9)
    _run(lp, vdate, size=9)
    if policy.approval_date:
        ap = body[0].add_paragraph()
        _run(ap, "Approval date: ", bold=True, size=9)
        _run(ap, policy.approval_date, size=9)
    op = body[1].paragraphs[0]
    _run(op, "Owner: ", bold=True, size=9)
    _run(op, policy.owner or "—", size=9)
    rp = body[1].add_paragraph()
    _run(rp, "Approver: ", bold=True, size=9)
    _run(rp, policy.approver or "Executive Committee", size=9)


def build_docx(policy, out_path, pdf_path=None):
    doc = Document()
    # Base style
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(policy.body_size or 11)
    normal.font.color.rgb = NAVY

    # Cover: page 1 of the PDF as a full-bleed image when available, else a
    # styled text title page as a fallback.
    have_cover = bool(pdf_path) and _cover_image_page(doc, pdf_path)
    if not have_cover:
        _normal_margins(doc.sections[0])
        for _ in range(3):
            doc.add_paragraph()
        _run(doc.add_paragraph(), policy.title, bold=True, size=30)
        if policy.cover_year and policy.year:
            _run(doc.add_paragraph(), str(policy.year), bold=True, size=20, color=GREY)

    # New section for the body, with the running header/footer furniture.
    doc.add_section(WD_SECTION.NEW_PAGE)
    body_sec = doc.sections[-1]
    _normal_margins(body_sec)
    _content_furniture(body_sec, policy)

    if policy.lead_title:
        _heading(doc, policy.title, 1)
    _render_blocks(doc, policy.blocks, size=policy.body_size or 11,
                    hanging_indent=getattr(policy, "hanging_indent", False))
    _version_card(doc, policy)
    doc.save(out_path)
    return out_path
